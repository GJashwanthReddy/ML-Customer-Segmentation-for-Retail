import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def clear_and_set_text(shape, paragraphs_data):
    """
    Completely clears all existing paragraphs from shape's text_frame,
    and populates it with paragraphs_data.
    paragraphs_data is a list of dicts:
    [{'text': str, 'font_name': str, 'size': float, 'bold': bool, 'italic': bool, 'color': (r,g,b), 'align': PP_ALIGN, 'space_after': float}]
    """
    tf = shape.text_frame
    tf.word_wrap = True
    
    # Keep the first paragraph, delete all remaining paragraphs from XML
    p0 = tf.paragraphs[0]
    p0.text = ""
    for p in list(tf.paragraphs)[1:]:
        p_elem = p._p
        p_elem.getparent().remove(p_elem)
        
    for idx, data in enumerate(paragraphs_data):
        p = p0 if idx == 0 else tf.add_paragraph()
        p.text = ""
        run = p.add_run()
        run.text = data.get('text', '')
        run.font.name = data.get('font_name', 'Calibri')
        run.font.size = Pt(data.get('size', 12.0))
        run.font.bold = data.get('bold', False)
        run.font.italic = data.get('italic', False)
        color = data.get('color', (0x1F, 0x24, 0x30))
        run.font.color.rgb = RGBColor(*color)
        if 'align' in data and data['align'] is not None:
            p.alignment = data['align']
        if 'space_after' in data:
            p.space_after = Pt(data['space_after'])

def build_presentation():
    prs = Presentation('Minor_Project_Model_PPT_Template_BACKUP.pptx')
    print(f"Loaded presentation with {len(prs.slides)} slides.")

    # Color Palette matching KL University template
    C_NAVY = (0x1B, 0x2A, 0x55)
    C_DARK_NAVY = (0x10, 0x1B, 0x3A)
    C_GOLD = (0xC8, 0xA0, 0x2C)
    C_CHARCOAL = (0x1F, 0x24, 0x30)
    C_CRIMSON = (0xA4, 0x16, 0x1A)
    C_GRAY = (0x6B, 0x72, 0x80)
    C_WHITE = (0xFF, 0xFF, 0xFF)
    C_LIGHT_BLUE = (0xAE, 0xBB, 0xD6)
    C_MUTED_BLUE = (0x8C, 0x9B, 0xBC)
    C_GREEN = (0x1E, 0x7A, 0x46)
    C_BOX_BG = (0xF1, 0xF4, 0xF9)
    C_BOX_BORDER = (0xDD, 0xE3, 0xEC)

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides[0]
    # Shape 2: Title
    clear_and_set_text(s1.shapes[2], [{
        'text': "Customer Segmentation for Retail",
        'font_name': 'Cambria',
        'size': 38.0,
        'bold': True,
        'color': C_WHITE
    }])
    # Shape 3: Subtitle
    clear_and_set_text(s1.shapes[3], [{
        'text': "Unsupervised Machine Learning for Retail Customer Behavioral Analysis",
        'font_name': 'Calibri',
        'size': 13.5,
        'bold': False,
        'color': C_GOLD
    }])
    # Shape 6: Students (Team size: 2)
    clear_and_set_text(s1.shapes[6], [
        {
            'text': "Student 1 Name   [Roll No. / ID]",
            'font_name': 'Calibri',
            'size': 14.5,
            'bold': False,
            'color': C_WHITE,
            'space_after': 4.0
        },
        {
            'text': "Student 2 Name   [Roll No. / ID]",
            'font_name': 'Calibri',
            'size': 14.5,
            'bold': False,
            'color': C_WHITE
        }
    ])
    # Shape 9: Guide name
    clear_and_set_text(s1.shapes[9], [{
        'text': "[Faculty Guide Name]",
        'font_name': 'Calibri',
        'size': 16.0,
        'bold': True,
        'color': C_WHITE
    }])
    # Shape 10: Guide designation
    clear_and_set_text(s1.shapes[10], [
        {
            'text': "Assistant Professor, Department of CSE",
            'font_name': 'Calibri',
            'size': 13.0,
            'bold': False,
            'color': C_LIGHT_BLUE
        },
        {
            'text': "KLH Bachupally Campus",
            'font_name': 'Calibri',
            'size': 13.0,
            'bold': False,
            'color': C_LIGHT_BLUE
        }
    ])
    # Shape 11: Footer metadata
    clear_and_set_text(s1.shapes[11], [{
        'text': "B.Tech II Year  |  Minor Project - Review 1  |  Academic Year 2026-27  |  KL University",
        'font_name': 'Calibri',
        'size': 11.5,
        'bold': False,
        'color': C_MUTED_BLUE
    }])

    # =========================================================================
    # SLIDE 2: Agenda (Preserve exactly)
    # =========================================================================
    # All 12 items already match perfectly with project structure!

    # =========================================================================
    # SLIDE 3: Introduction & Domain Background
    # =========================================================================
    s3 = prs.slides[2]
    # Shape 3: Left text card
    clear_and_set_text(s3.shapes[3], [
        {
            'text': "What is the domain?",
            'font_name': 'Calibri',
            'size': 14.5,
            'bold': True,
            'color': C_NAVY
        },
        {
            'text': "Modern retail and e-commerce enterprises generate massive volumes of customer transaction data daily. Customer purchasing behaviors vary substantially across frequency, timing, and expenditure. Treating all customers uniformly leads to suboptimal marketing efficiency and customer attrition.",
            'font_name': 'Calibri',
            'size': 12.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 6.0
        },
        {
            'text': "Why is it relevant today?",
            'font_name': 'Calibri',
            'size': 14.5,
            'bold': True,
            'color': C_NAVY
        },
        {
            'text': "Studies in retail analytics indicate that acquiring a new customer costs up to 5 times more than retaining an existing customer [6]. Automated customer segmentation enables personalized marketing campaigns, improves retention rates, and maximizes customer lifetime value.",
            'font_name': 'Calibri',
            'size': 12.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 6.0
        },
        {
            'text': "Where does our project sit?",
            'font_name': 'Calibri',
            'size': 14.5,
            'bold': True,
            'color': C_NAVY
        },
        {
            'text': "Our project applies unsupervised machine learning to retail transaction logs to discover natural customer cohorts using Recency, Frequency, and Monetary (RFM) behavioral feature extraction and clustering algorithms.",
            'font_name': 'Calibri',
            'size': 12.0,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])

    # Shape 5: Right card concept illustration
    clear_and_set_text(s3.shapes[5], [
        {
            'text': "RFM BEHAVIORAL MODEL",
            'font_name': 'Calibri',
            'size': 13.0,
            'bold': True,
            'color': C_NAVY,
            'space_after': 4.0
        },
        {
            'text': "• Recency (R): Days since last purchase; distinguishes active from churning buyers.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 3.0
        },
        {
            'text': "• Frequency (F): Total count of purchases; measures brand engagement and loyalty.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 3.0
        },
        {
            'text': "• Monetary (M): Cumulative amount spent; identifies top revenue-generating customers.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 6.0
        },
        {
            'text': "Unsupervised Clustering discovers organic customer segments directly from these behavioral vectors without predefined labels.",
            'font_name': 'Calibri',
            'size': 10.5,
            'italic': True,
            'color': C_NAVY
        }
    ])

    # Shape 7: Panel tip header
    clear_and_set_text(s3.shapes[7], [{
        'text': "Project Grounding",
        'font_name': 'Calibri',
        'size': 10.5,
        'bold': True,
        'color': C_CRIMSON
    }])
    # Shape 8: Panel tip text
    clear_and_set_text(s3.shapes[8], [{
        'text': "All domain insights and customer retention benchmarks are grounded in peer-reviewed retail analytics literature [1], [6].",
        'font_name': 'Calibri',
        'size': 11.5,
        'bold': False,
        'color': C_CHARCOAL
    }])

    # =========================================================================
    # SLIDE 4: Problem Statement & Motivation
    # =========================================================================
    s4 = prs.slides[3]
    # Shape 4: Heading
    clear_and_set_text(s4.shapes[4], [{
        'text': "PROBLEM STATEMENT",
        'font_name': 'Calibri',
        'size': 10.5,
        'bold': True,
        'color': C_CRIMSON
    }])
    # Shape 5: Statement
    clear_and_set_text(s4.shapes[5], [{
        'text': "\"Retail businesses often struggle to identify distinct customer groups from large, high-dimensional transaction datasets, which limits their ability to design targeted marketing and customer retention strategies.\"",
        'font_name': 'Cambria',
        'size': 17.5,
        'bold': True,
        'color': C_NAVY
    }])
    # Card 1 (Shape 9 & 10)
    clear_and_set_text(s4.shapes[9], [{
        'text': "Who is affected?",
        'font_name': 'Calibri',
        'size': 14.5,
        'bold': True,
        'color': C_NAVY
    }])
    clear_and_set_text(s4.shapes[10], [{
        'text': "Retail enterprises, e-commerce businesses, CRM managers, and marketing strategists seeking data-driven customer retention and personalized engagement.",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Card 2 (Shape 14 & 15)
    clear_and_set_text(s4.shapes[14], [{
        'text': "What is the cost of the problem?",
        'font_name': 'Calibri',
        'size': 14.5,
        'bold': True,
        'color': C_NAVY
    }])
    clear_and_set_text(s4.shapes[15], [{
        'text': "Generic one-size-fits-all marketing wastes advertising budgets, fails to prevent customer churn among at-risk buyers, and misses opportunities to reward high-value loyal customers.",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Card 3 (Shape 19 & 20)
    clear_and_set_text(s4.shapes[19], [{
        'text': "Why solve it now?",
        'font_name': 'Calibri',
        'size': 14.5,
        'bold': True,
        'color': C_NAVY
    }])
    clear_and_set_text(s4.shapes[20], [{
        'text': "Transaction volumes make manual segmentation impossible; unsupervised ML automatically uncovers customer behavioral cohorts directly from transaction logs without human bias.",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': False,
        'color': C_CHARCOAL
    }])

    # =========================================================================
    # SLIDE 5: Literature Survey
    # =========================================================================
    s5 = prs.slides[4]
    # Subtitle (Shape 2)
    clear_and_set_text(s5.shapes[2], [{
        'text': "Comparative analysis of 4 key recent academic papers (2020–2024) establishing the foundation for retail customer segmentation.",
        'font_name': 'Calibri',
        'size': 10.5,
        'bold': False,
        'color': C_GRAY
    }])

    # Table (Shape 3)
    table_shape = s5.shapes[3]
    t = table_shape.table

    papers_data = [
        # Paper 1
        ("1",
         "Wani et al. (2023) [1]",
         "Unleashing Customer Insights: Segmentation Through Machine Learning (K-Means on RFM)",
         "Retail Transaction Data / Python, Scikit-learn",
         "Successfully segmented retail customers into behavioral cohorts; optimized K using Elbow method to boost marketing efficiency.",
         "Relied strictly on standard K-Means without comparative evaluation against hierarchical algorithms; minimal 2D visual projection."),
        # Paper 2
        ("2",
         "Teslenko et al. (2023) [2]",
         "Comparative Analysis of Five Clustering Algorithms for Market Segmentation",
         "E-Commerce Customer Records / Scikit-learn, NumPy",
         "Benchmarked K-Means, Agglomerative, and DBSCAN; showed K-Means speed and Agglomerative hierarchy depth for customer groupings.",
         "Focused purely on mathematical clustering metrics without RFM feature extraction or translating clusters into retail business personas."),
        # Paper 3
        ("3",
         "Li et al. (2022) [3]",
         "E-commerce Customer Segmentation based on RFM Weighted K-means",
         "E-commerce Transaction Logs / Python, Scikit-learn",
         "Weighting RFM features improved cluster compactness and differentiated high-spend customer tiers over raw transactional attributes.",
         "Did not incorporate PCA dimensionality reduction to visualize high-dimensional customer boundaries for non-technical stakeholders."),
        # Paper 4
        ("4",
         "Akande et al. (2024) [4]",
         "Customer Segmentation through RFM Analysis and K-means Clustering",
         "UCI Online Retail Dataset / Python, Pandas, Seaborn",
         "Derived RFM metrics from raw retail transactions; validated cluster partitions using Elbow inertia and Silhouette analysis.",
         "Examined K-Means in isolation without hierarchical dendrogram validation, and lacked a systematic framework for business persona profiling.")
    ]

    for row_idx, p_data in enumerate(papers_data):
        r = row_idx + 1
        for col_idx in range(6):
            cell = t.cell(r, col_idx)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            for p_rem in list(tf.paragraphs)[1:]:
                p_elem = p_rem._p
                p_elem.getparent().remove(p_elem)
            run = p.add_run()
            run.text = p_data[col_idx]
            run.font.name = "Calibri"
            run.font.size = Pt(9.5 if col_idx in [2, 4, 5] else 10.0)
            run.font.bold = (col_idx == 0)
            run.font.color.rgb = RGBColor(*C_CHARCOAL)

    # Footnote (Shape 4)
    clear_and_set_text(s5.shapes[4], [{
        'text': "All identified limitations directly inform the research gaps addressed in Section 4 (Additional literature studied: [5]–[9]).",
        'font_name': 'Calibri',
        'size': 11.5,
        'bold': False,
        'color': C_CHARCOAL
    }])

    # =========================================================================
    # SLIDE 6: Research Gap Identified
    # =========================================================================
    s6 = prs.slides[5]
    # Card 1: GAP 1
    clear_and_set_text(s6.shapes[3], [{
        'text': "GAP 1: DATA & FEATURE ENGINEERING",
        'font_name': 'Calibri',
        'size': 10.5,
        'bold': True,
        'color': C_CRIMSON
    }])
    clear_and_set_text(s6.shapes[4], [{
        'text': "Existing studies often cluster raw transaction summaries or unweighted features without standardizing skewed RFM behavioral distributions, causing centroid bias toward extreme monetary values (Li et al., 2022; Akande et al., 2024).",
        'font_name': 'Calibri',
        'size': 11.5,
        'bold': False,
        'color': C_CHARCOAL
    }])

    # Card 2: GAP 2
    clear_and_set_text(s6.shapes[6], [{
        'text': "GAP 2: METHODOLOGY & COMPARISON",
        'font_name': 'Calibri',
        'size': 10.5,
        'bold': True,
        'color': C_CRIMSON
    }])
    clear_and_set_text(s6.shapes[7], [{
        'text': "Most retail segmentation models rely exclusively on standard K-Means without comparative evaluation against hierarchical clustering, leaving cluster stability, tree hierarchies, and sensitivity to centroid initialization unverified (Wani et al., 2023).",
        'font_name': 'Calibri',
        'size': 11.5,
        'bold': False,
        'color': C_CHARCOAL
    }])

    # Card 3: GAP 3
    clear_and_set_text(s6.shapes[9], [{
        'text': "GAP 3: VISUALIZATION & INTERPRETATION",
        'font_name': 'Calibri',
        'size': 10.5,
        'bold': True,
        'color': C_CRIMSON
    }])
    clear_and_set_text(s6.shapes[10], [{
        'text': "Prior works output mathematical cluster labels but lack 2D PCA dimensionality reduction visualization and structured mapping into actionable retail customer personas for business decision-making (Teslenko et al., 2023; Li et al., 2022).",
        'font_name': 'Calibri',
        'size': 11.5,
        'bold': False,
        'color': C_CHARCOAL
    }])

    # Right Card: OUR PROJECT ADDRESSES
    clear_and_set_text(s6.shapes[12], [{
        'text': "OUR PROJECT ADDRESSES",
        'font_name': 'Calibri',
        'size': 10.5,
        'bold': True,
        'color': C_GOLD
    }])
    clear_and_set_text(s6.shapes[13], [{
        'text': "Our project addresses all three gaps by building an integrated unsupervised pipeline: extracting and scaling RFM features, benchmarking K-Means against Hierarchical Clustering with Elbow and Silhouette metrics, applying PCA for 2D visualization, and formulating data-driven retail customer personas.",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': False,
        'color': C_NAVY
    }])
    clear_and_set_text(s6.shapes[14], [{
        'text': "Traceability: Every identified gap is directly grounded in the limitations column of our Literature Survey (Slide 5).",
        'font_name': 'Calibri',
        'size': 11.0,
        'italic': True,
        'color': C_CHARCOAL
    }])
    clear_and_set_text(s6.shapes[15], [{
        'text': "Methodological Impact: Combines RFM feature engineering with multi-algorithm clustering and PCA visual interpretation.",
        'font_name': 'Calibri',
        'size': 11.5,
        'bold': False,
        'color': C_CHARCOAL
    }])

    # =========================================================================
    # SLIDE 7: Objectives of the Project
    # =========================================================================
    s7 = prs.slides[6]
    # Subtitle
    clear_and_set_text(s7.shapes[2], [{
        'text': "Four action-oriented, measurable objectives formulated for Minor Project Review 1.",
        'font_name': 'Calibri',
        'size': 10.5,
        'bold': False,
        'color': C_GRAY
    }])
    # Objective 1
    clear_and_set_text(s7.shapes[6], [{
        'text': "To collect and preprocess retail transaction data suitable for customer segmentation by cleaning missing customer IDs, removing cancelled/invalid transactions, and standardizing date formats.",
        'font_name': 'Calibri',
        'size': 13.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Objective 2
    clear_and_set_text(s7.shapes[10], [{
        'text': "To extract Recency, Frequency, and Monetary (RFM) behavioral features from customer transaction records and apply feature scaling to ensure uniform attribute contribution.",
        'font_name': 'Calibri',
        'size': 13.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Objective 3
    clear_and_set_text(s7.shapes[14], [{
        'text': "To implement and compare clustering approaches, primarily K-Means and Hierarchical Clustering, using the Elbow Method and Silhouette Score to determine optimal cluster configuration.",
        'font_name': 'Calibri',
        'size': 13.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Objective 4
    clear_and_set_text(s7.shapes[18], [{
        'text': "To visualize and interpret the resulting customer segments using Principal Component Analysis (PCA) and derive meaningful retail business personas and actionable marketing insights.",
        'font_name': 'Calibri',
        'size': 13.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Tip
    clear_and_set_text(s7.shapes[19], [{
        'text': "Review 1 Alignment: Objectives will be verified one-to-one against experimental deliverables in Review 2 and Final Review.",
        'font_name': 'Calibri',
        'size': 11.5,
        'italic': True,
        'color': C_GRAY
    }])

    # =========================================================================
    # SLIDE 8: Proposed System Architecture
    # =========================================================================
    s8 = prs.slides[7]
    # Update top row 4 overview cards
    clear_and_set_text(s8.shapes[4], [
        {
            'text': "• UCI Online Retail Dataset",
            'font_name': 'Calibri',
            'size': 11.5,
            'bold': True,
            'color': C_NAVY
        },
        {
            'text': "• Customer IDs & transaction logs\n• Invoices, quantities, unit prices",
            'font_name': 'Calibri',
            'size': 10.5,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])
    clear_and_set_text(s8.shapes[8], [
        {
            'text': "• Data Cleaning & Filtering",
            'font_name': 'Calibri',
            'size': 11.5,
            'bold': True,
            'color': C_NAVY
        },
        {
            'text': "• Filter returns & missing IDs\n• RFM Calculation & StandardScaler",
            'font_name': 'Calibri',
            'size': 10.5,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])
    clear_and_set_text(s8.shapes[12], [
        {
            'text': "• K-Means Clustering",
            'font_name': 'Calibri',
            'size': 11.5,
            'bold': True,
            'color': C_NAVY
        },
        {
            'text': "• Elbow & Silhouette evaluation\n• Hierarchical clustering validation",
            'font_name': 'Calibri',
            'size': 10.5,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])
    clear_and_set_text(s8.shapes[16], [
        {
            'text': "• PCA 2D Visual Projections",
            'font_name': 'Calibri',
            'size': 11.5,
            'bold': True,
            'color': C_NAVY
        },
        {
            'text': "• RFM cluster centroid profiles\n• Retail marketing & retention plans",
            'font_name': 'Calibri',
            'size': 10.5,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])

    # Remove placeholder shapes 17, 18, 19, 20, 21 from Slide 8
    shapes_to_remove = [s8.shapes[21], s8.shapes[20], s8.shapes[19], s8.shapes[18], s8.shapes[17]]
    for sp in shapes_to_remove:
        sp_elem = sp._element
        sp_elem.getparent().remove(sp_elem)

    # Build the 8-stage S-flow diagram using PowerPoint shapes
    # Row 1 (Left to Right: Steps 1, 2, 3, 4): top = 4.05", height = 1.10"
    row1_steps = [
        ("Step 1: Retail Data", "UCI Online Retail logs", "Customer transaction records"),
        ("Step 2: Preprocessing", "Data cleaning & filtering", "Handle missing IDs & returns"),
        ("Step 3: RFM Extraction", "Recency, Frequency, Monetary", "Customer-level behavioral vectors"),
        ("Step 4: Feature Scaling", "StandardScaler normalization", "Uniform feature contribution")
    ]
    row1_arrow_labels = ["Raw CSV", "Clean Data", "RFM Vectors"]

    # Row 2 (Right to Left: Steps 5, 6, 7, 8): top = 5.45", height = 1.10"
    # Step 5 is directly under Step 4!
    # Step 5: Cluster Selection (Elbow + Silhouette)
    # Step 6: Clustering Models (K-Means & Hierarchical)
    # Step 7: PCA Visualization (Dimensionality Reduction)
    # Step 8: Customer Personas & Business Insights
    row2_steps = [
        ("Step 8: Retail Personas", "Cluster profiling & centroids", "Actionable business strategies"),
        ("Step 7: PCA Visualization", "Dimensionality reduction", "2D projection for visual separation"),
        ("Step 6: Clustering Models", "K-Means & Hierarchical", "Partition customers & validate tree"),
        ("Step 5: Cluster Evaluation", "Elbow Method & Silhouette", "Determine optimal cluster count K")
    ]
    row2_arrow_labels = ["Target Insights", "2D Projection", "Cluster IDs"]

    box_w = Inches(2.25)
    box_h = Inches(1.10)
    gap_x = Inches(1.03)

    # Add Row 1 (Steps 1 -> 2 -> 3 -> 4)
    for i in range(4):
        bx = Inches(0.60) + i * (box_w + gap_x)
        by = Inches(4.05)
        shape = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, box_w, box_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*C_BOX_BG)
        shape.line.color.rgb = RGBColor(*C_NAVY)
        shape.line.width = Pt(1.25)
        clear_and_set_text(shape, [
            {'text': row1_steps[i][0], 'font_name': 'Calibri', 'size': 10.5, 'bold': True, 'color': C_NAVY, 'align': PP_ALIGN.CENTER},
            {'text': row1_steps[i][1], 'font_name': 'Calibri', 'size': 9.0, 'bold': False, 'color': C_CHARCOAL, 'align': PP_ALIGN.CENTER},
            {'text': row1_steps[i][2], 'font_name': 'Calibri', 'size': 8.5, 'italic': True, 'color': C_GRAY, 'align': PP_ALIGN.CENTER}
        ])

        # Right Arrow after box
        if i < 3:
            ax = bx + box_w + Inches(0.12)
            ay = by + Inches(0.48)
            aw = Inches(0.79)
            ah = Inches(0.18)
            arrow = s8.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, ay, aw, ah)
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(*C_GOLD)
            arrow.line.fill.background()
            tx = s8.shapes.add_textbox(ax - Inches(0.08), by + Inches(0.15), aw + Inches(0.16), Inches(0.30))
            clear_and_set_text(tx, [{
                'text': row1_arrow_labels[i],
                'font_name': 'Calibri',
                'size': 8.5,
                'bold': True,
                'color': C_CRIMSON,
                'align': PP_ALIGN.CENTER
            }])

    # Vertical down connector between Step 4 and Step 5 on the right!
    # Step 4 is at left = Inches(0.60) + 3 * (2.25 + 1.03) = Inches(10.44)
    step4_bx = Inches(0.60) + 3 * (box_w + gap_x)
    down_ax = step4_bx + Inches(0.95)
    down_ay = Inches(5.17)
    down_arrow = s8.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, down_ax, down_ay, Inches(0.35), Inches(0.24))
    down_arrow.fill.solid()
    down_arrow.fill.fore_color.rgb = RGBColor(*C_GOLD)
    down_arrow.line.fill.background()

    down_label = s8.shapes.add_textbox(step4_bx - Inches(0.40), Inches(5.15), Inches(1.30), Inches(0.26))
    clear_and_set_text(down_label, [{
        'text': "Scaled Matrix",
        'font_name': 'Calibri',
        'size': 8.5,
        'bold': True,
        'color': C_CRIMSON,
        'align': PP_ALIGN.RIGHT
    }])

    # Add Row 2 (Steps 8 <- 7 <- 6 <- 5)
    # i=0 is Step 8 at left, i=3 is Step 5 at right
    for i in range(4):
        bx = Inches(0.60) + i * (box_w + gap_x)
        by = Inches(5.45)
        shape = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, box_w, box_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*C_BOX_BG)
        shape.line.color.rgb = RGBColor(*C_NAVY)
        shape.line.width = Pt(1.25)
        clear_and_set_text(shape, [
            {'text': row2_steps[i][0], 'font_name': 'Calibri', 'size': 10.5, 'bold': True, 'color': C_NAVY, 'align': PP_ALIGN.CENTER},
            {'text': row2_steps[i][1], 'font_name': 'Calibri', 'size': 9.0, 'bold': False, 'color': C_CHARCOAL, 'align': PP_ALIGN.CENTER},
            {'text': row2_steps[i][2], 'font_name': 'Calibri', 'size': 8.5, 'italic': True, 'color': C_GRAY, 'align': PP_ALIGN.CENTER}
        ])

        # Left Arrow pointing from right box to left box
        if i < 3:
            ax = bx + box_w + Inches(0.12)
            ay = by + Inches(0.48)
            aw = Inches(0.79)
            ah = Inches(0.18)
            arrow = s8.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, ax, ay, aw, ah)
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(*C_GOLD)
            arrow.line.fill.background()
            tx = s8.shapes.add_textbox(ax - Inches(0.08), by + Inches(0.15), aw + Inches(0.16), Inches(0.30))
            clear_and_set_text(tx, [{
                'text': row2_arrow_labels[i],
                'font_name': 'Calibri',
                'size': 8.5,
                'bold': True,
                'color': C_CRIMSON,
                'align': PP_ALIGN.CENTER
            }])

    # =========================================================================
    # SLIDE 9: Methodology & Modules
    # =========================================================================
    s9 = prs.slides[8]
    # Module 1 (Shape 5 & 6)
    clear_and_set_text(s9.shapes[5], [{
        'text': "Module 1 – Data Collection & Preprocessing",
        'font_name': 'Calibri',
        'size': 14.5,
        'bold': True,
        'color': C_NAVY
    }])
    clear_and_set_text(s9.shapes[6], [
        {
            'text': "• Planned Source: UCI Online Retail Dataset containing transactional records.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Data Cleaning: Filter records with missing CustomerID and invalid quantities.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Handling Returns: Exclude cancelled transactions (invoice codes starting with 'C').",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Transaction Value: Compute line total = Quantity × UnitPrice; parse datetime stamps.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])

    # Module 2 (Shape 10 & 11)
    clear_and_set_text(s9.shapes[10], [{
        'text': "Module 2 – RFM Feature Engineering",
        'font_name': 'Calibri',
        'size': 14.5,
        'bold': True,
        'color': C_NAVY
    }])
    clear_and_set_text(s9.shapes[11], [
        {
            'text': "• Recency (R): Calculate days elapsed between latest purchase and reference date.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Frequency (F): Count total unique purchase transactions per customer.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Monetary (M): Aggregate total amount spent across all customer invoices.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Feature Scaling: Apply log-transformation & StandardScaler to normalize distributions.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])

    # Module 3 (Shape 15 & 16)
    clear_and_set_text(s9.shapes[15], [{
        'text': "Module 3 – Clustering & Comparative Evaluation",
        'font_name': 'Calibri',
        'size': 14.5,
        'bold': True,
        'color': C_NAVY
    }])
    clear_and_set_text(s9.shapes[16], [
        {
            'text': "• K-Means Clustering: Partition customer vectors into K clusters using Euclidean distance.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Optimal K Selection: Evaluate within-cluster inertia (Elbow) and Silhouette Scores.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Hierarchical Clustering: Apply Agglomerative Clustering with Ward's linkage for comparison.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Dimensionality Reduction: Apply PCA to project 3D RFM features into 2D space.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])

    # Module 4 (Shape 20 & 21)
    clear_and_set_text(s9.shapes[20], [{
        'text': "Module 4 – Customer Profiling & Business Insights",
        'font_name': 'Calibri',
        'size': 14.5,
        'bold': True,
        'color': C_NAVY
    }])
    clear_and_set_text(s9.shapes[21], [
        {
            'text': "• Cluster Profiling: Compute mean and median RFM characteristics for each cluster.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Persona Mapping: Assign empirical descriptions (e.g., Loyal, Frequent, Occasional, At-Risk).",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Retail Strategies: Formulate targeted retention, reward, and reactivation recommendations.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 2.0
        },
        {
            'text': "• Interface (Planned): Interactive Streamlit dashboard for business exploration.",
            'font_name': 'Calibri',
            'size': 11.0,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])

    # =========================================================================
    # SLIDE 10: Tools and Technologies
    # =========================================================================
    s10 = prs.slides[9]
    # Software (Shape 4)
    clear_and_set_text(s10.shapes[4], [
        {
            'text': "Language: Python 3.10+",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "Data Processing: Pandas, NumPy",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "Machine Learning: Scikit-learn (K-Means, Agglomerative, PCA)",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "Visualization: Matplotlib, Seaborn",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "User Interface: Streamlit (Planned for prototype deployment)",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "IDE / Environment: VS Code / Jupyter Notebook / Google Colab",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])

    # Hardware (Shape 7)
    clear_and_set_text(s10.shapes[7], [
        {
            'text': "Processor: Intel Core i5 / AMD Ryzen 5 or equivalent",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "RAM: 8 GB minimum (16 GB recommended for pandas efficiency)",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "Storage: 512 GB SSD",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "Compute / GPU: Standard CPU compute (GPU not required)",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "Operating System: Windows 11 / Linux",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "Version Control: Git & GitHub repository",
            'font_name': 'Calibri',
            'size': 13.5,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])

    # =========================================================================
    # SLIDE 11: Work Completed / Demo
    # =========================================================================
    s11 = prs.slides[10]
    # Update Status Items
    # Row 1: Literature survey (Shape 4 & 5) -> Completed (Green)
    clear_and_set_text(s11.shapes[4], [{
        'text': "Literature survey & domain study",
        'font_name': 'Calibri',
        'size': 13.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Row 2: Dataset selection (Shape 8 & 9) -> Completed (Green)
    clear_and_set_text(s11.shapes[8], [{
        'text': "Dataset identification (UCI Retail)",
        'font_name': 'Calibri',
        'size': 13.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Row 3: Pre-processing pipeline (Shape 12 & 13) -> In progress
    clear_and_set_text(s11.shapes[12], [{
        'text': "Data preprocessing pipeline",
        'font_name': 'Calibri',
        'size': 13.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Row 4: RFM Feature Engineering (Shape 16 & 17) -> In progress
    clear_and_set_text(s11.shapes[16], [{
        'text': "RFM feature engineering logic",
        'font_name': 'Calibri',
        'size': 13.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Row 5: Clustering models (Shape 20 & 21) -> In progress / To be evaluated
    clear_and_set_text(s11.shapes[20], [{
        'text': "Clustering models (K-Means/Hierarchical)",
        'font_name': 'Calibri',
        'size': 13.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    clear_and_set_text(s11.shapes[21], [{
        'text': "To be evaluated",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': True,
        'color': C_GOLD
    }])
    # Update circle for Row 5 to Gold
    s11.shapes[19].fill.solid()
    s11.shapes[19].fill.fore_color.rgb = RGBColor(*C_GOLD)

    # Row 6: Testing & evaluation (Shape 24 & 25) -> Planned
    clear_and_set_text(s11.shapes[24], [{
        'text': "PCA visualization & UI dashboard",
        'font_name': 'Calibri',
        'size': 13.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    clear_and_set_text(s11.shapes[25], [{
        'text': "Planned",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': True,
        'color': C_GRAY
    }])

    # Right side: Shape 27 (Replace screenshot placeholder with structured progress card)
    clear_and_set_text(s11.shapes[27], [
        {
            'text': "CURRENT IMPLEMENTATION FOCUS: DATA PIPELINE & RFM SCHEMA",
            'font_name': 'Calibri',
            'size': 11.5,
            'bold': True,
            'color': C_NAVY,
            'space_after': 4.0
        },
        {
            'text': "Planned Input Schema (UCI Online Retail Dataset):",
            'font_name': 'Calibri',
            'size': 10.5,
            'bold': True,
            'color': C_CHARCOAL
        },
        {
            'text': "• Attributes: InvoiceNo, StockCode, Description, Quantity, InvoiceDate, UnitPrice, CustomerID, Country\n• Format: Transaction-level tabular records (~541k instances in raw corpus)",
            'font_name': 'Calibri',
            'size': 9.5,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 4.0
        },
        {
            'text': "Mathematical RFM Extraction Formulae:",
            'font_name': 'Calibri',
            'size': 10.5,
            'bold': True,
            'color': C_CHARCOAL
        },
        {
            'text': "• Recency (R) = max(InvoiceDate) - max(InvoiceDate_i)  [days elapsed]\n• Frequency (F) = Count(Unique InvoiceNo per CustomerID)\n• Monetary (M) = Sigma (Quantity x UnitPrice) per CustomerID",
            'font_name': 'Calibri',
            'size': 9.5,
            'bold': False,
            'color': C_NAVY,
            'space_after': 4.0
        },
        {
            'text': "Active Data Cleaning Rules:",
            'font_name': 'Calibri',
            'size': 10.5,
            'bold': True,
            'color': C_CHARCOAL
        },
        {
            'text': "• Dropping null CustomerID records; removing negative quantities (cancellations/returns).\n• Applying log-transform and StandardScaler to eliminate heavy monetary skewness.",
            'font_name': 'Calibri',
            'size': 9.5,
            'bold': False,
            'color': C_CHARCOAL
        }
    ])

    # Shape 29: Bottom tip
    clear_and_set_text(s11.shapes[29], [{
        'text': "Review 1 Status: Data ingestion & RFM engineering logic are formulated. Model clustering and evaluation will execute once normalized vectors are finalized.",
        'font_name': 'Calibri',
        'size': 11.5,
        'bold': False,
        'color': C_CHARCOAL
    }])

    # =========================================================================
    # SLIDE 12: Timeline & Review Plan
    # =========================================================================
    s12 = prs.slides[11]
    # Review 0 (Week 2)
    clear_and_set_text(s12.shapes[6], [{
        'text': "Domain selection, team finalization & abstract approved",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Review 1 (Week 5 - Current)
    clear_and_set_text(s12.shapes[9], [{
        'text': "Week 5 (Current)",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': True,
        'color': C_GOLD
    }])
    clear_and_set_text(s12.shapes[10], [{
        'text': "Literature survey, research gaps, architecture & RFM methodology",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Review 2 (Week 9)
    clear_and_set_text(s12.shapes[14], [{
        'text': "Data preprocessing, RFM scaling, K-Means & Hierarchical evaluation",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Final Review (Week 14)
    clear_and_set_text(s12.shapes[18], [{
        'text': "PCA visualizations, customer profiling, dashboard UI & report",
        'font_name': 'Calibri',
        'size': 12.0,
        'bold': False,
        'color': C_CHARCOAL
    }])
    # Shape 19: Note
    clear_and_set_text(s12.shapes[19], [{
        'text': "Status Tracking: Review 0 and Review 1 completed on schedule; data engineering progressing toward Review 2.",
        'font_name': 'Calibri',
        'size': 11.5,
        'bold': False,
        'color': C_GRAY
    }])
    # Shape 21: Panel tip
    clear_and_set_text(s12.shapes[21], [{
        'text': "Academic Compliance: All Review 1 deliverables achieved honestly without fabricated results; adequate timeline remains to complete subsequent clustering stages.",
        'font_name': 'Calibri',
        'size': 12.5,
        'bold': True,
        'color': C_CHARCOAL
    }])

    # =========================================================================
    # SLIDE 13: Expected Outcome & Conclusion
    # =========================================================================
    s13 = prs.slides[12]
    # Deliverable (Shape 4)
    clear_and_set_text(s13.shapes[4], [{
        'text': "A functional customer segmentation system that processes retail transaction logs, computes normalized RFM features, discovers behavioral customer segments, and presents 2D PCA cluster visualizations.",
        'font_name': 'Cambria',
        'size': 13.5,
        'bold': True,
        'color': C_NAVY
    }])
    # Measurable Target (Shape 7)
    clear_and_set_text(s13.shapes[7], [{
        'text': "Evaluation of cluster cohesion via Silhouette Score and Elbow inertia curve to identify optimal K, avoiding arbitrary or subjective segment definitions.",
        'font_name': 'Cambria',
        'size': 13.5,
        'bold': True,
        'color': C_WHITE
    }])
    # Beyond the lab (Shape 10)
    clear_and_set_text(s13.shapes[10], [{
        'text': "Retailers and marketing teams can utilize discovered customer personas to optimize promotional budgets, design loyalty rewards, and deploy re-engagement campaigns.",
        'font_name': 'Cambria',
        'size': 13.5,
        'bold': True,
        'color': C_NAVY
    }])
    # Conclusion (Shape 13)
    clear_and_set_text(s13.shapes[13], [
        {
            'text': "• Conclusion: The project establishes a structured unsupervised learning framework to transform raw retail sales data into actionable customer intelligence using RFM analysis, K-Means, and Hierarchical clustering.",
            'font_name': 'Calibri',
            'size': 12.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 3.0
        },
        {
            'text': "• Final Deliverable: By the final review, the system will output validated cluster profiles, PCA visualizations, and targeted business strategy mappings.",
            'font_name': 'Calibri',
            'size': 12.0,
            'bold': False,
            'color': C_CHARCOAL,
            'space_after': 3.0
        },
        {
            'text': "• Future Scope: Real-time streaming segmentation on live POS data and automated enterprise CRM integration are intentionally reserved for future work.",
            'font_name': 'Calibri',
            'size': 12.0,
            'bold': False,
            'color': C_CRIMSON
        }
    ])

    # =========================================================================
    # SLIDE 14: References
    # =========================================================================
    s14 = prs.slides[13]
    # Subtitle (Shape 2)
    clear_and_set_text(s14.shapes[2], [{
        'text': "IEEE style, numbered in citation order. Peer-reviewed research papers and dataset repository for Minor Project Review 1.",
        'font_name': 'Calibri',
        'size': 10.5,
        'bold': False,
        'color': C_GRAY
    }])
    # References list (Shape 4)
    refs = [
        '[1]  A. Wani, M. Priyanka, and R. Prasath, "Unleashing Customer Insights: Segmentation Through Machine Learning," in Proc. World Conf. on Communication & Computing (WCONF), Raipur, India, 2023, pp. 1–6.',
        '[2]  D. Teslenko, A. Sorokina, K. Smelyakov, and O. Filipov, "Comparative Analysis of the Applicability of Five Clustering Algorithms for Market Segmentation," in Proc. IEEE Open Conf. of Electrical, Electronic and Information Sciences (eStream), Vilnius, Lithuania, 2023, pp. 1–5.',
        '[3]  P. Li, C. Wang, J. Wu, and R. Madleňák, "An E-commerce Customer Segmentation Method based on RFM Weighted K-means," in Proc. Int. Conf. on Management Engineering, Software Engineering and Service Sciences (ICMSS), Wuhan, China, 2022, pp. 61–68.',
        '[4]  O. Akande, A. O. Badmus, and B. S. Adekunle, "Customer Segmentation through RFM Analysis and K-means Clustering: Leveraging Data-Driven Insights for Effective Marketing Strategy," Ceddi J. Inf. Syst. Technol., vol. 2, no. 1, pp. 12–21, 2024.',
        '[5]  Y. Huang, M. Zhang, and Y. He, "Research on Improved RFM Customer Segmentation Model Based on K-Means Algorithm," in Proc. IEEE Int. Conf. on Artificial Intelligence and Computer Applications (ICAICA), Dalian, China, 2020, pp. 1010–1014.',
        '[6]  D. Chen, S. L. Sain, and K. Guo, "Data mining for the online retail industry: A case study of RFM model-based customer segmentation using data mining," J. Database Mark. Cust. Strategy Manag., vol. 19, no. 3, pp. 197–208, 2012.',
        '[7]  D. Dua and C. Graff, "Online Retail Data Set," UCI Machine Learning Repository, University of California, Irvine, School of Information and Computer Sciences, 2019. [Online]. Available: https://archive.ics.uci.edu/dataset/352/online+retail',
        '[8]  P. J. Rousseeuw, "Silhouettes: A graphical aid to the interpretation and validation of cluster analysis," J. Comput. Appl. Math., vol. 20, pp. 53–65, 1987.',
        '[9]  J. MacQueen, "Some methods for classification and analysis of multivariate observations," in Proc. 5th Berkeley Symp. Math. Statist. Prob., vol. 1, 1967, pp. 281–297.'
    ]
    ref_paras = [{
        'text': r_text,
        'font_name': 'Calibri',
        'size': 9.5,
        'bold': False,
        'color': C_CHARCOAL,
        'space_after': 2.0
    } for r_text in refs]
    clear_and_set_text(s14.shapes[4], ref_paras)

    # Footnote (Shape 5)
    clear_and_set_text(s14.shapes[5], [{
        'text': "Validation: All listed references are cited directly in the Literature Survey (Slide 5), Domain Context (Slide 3), and Methodology.",
        'font_name': 'Calibri',
        'size': 11.5,
        'bold': True,
        'color': C_CRIMSON
    }])

    # =========================================================================
    # SLIDE 15: Thank You Slide
    # =========================================================================
    s15 = prs.slides[14]
    # Shape 3: Team / Guide / Dept
    clear_and_set_text(s15.shapes[3], [{
        'text': "Team: [Student 1 & Student 2]   |   Guide: [Faculty Guide Name]   |   Dept. of CSE, KLH Bachupally",
        'font_name': 'Calibri',
        'size': 14.0,
        'bold': False,
        'color': C_LIGHT_BLUE
    }])

    # Save to both target locations
    output_filename1 = "Customer_Segmentation_Minor_Project_Review_1.pptx"
    output_filename2 = "Minor_Project_Model_PPT_Template.pptx"
    
    prs.save(output_filename1)
    prs.save(output_filename2)
    print(f"Saved successfully to '{output_filename1}' and '{output_filename2}'!")

if __name__ == "__main__":
    build_presentation()
